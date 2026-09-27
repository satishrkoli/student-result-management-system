import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os, hashlib

FILE = "student_management.xlsx"

def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()

def create_file():
    if not os.path.exists(FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Students"
        ws.append(["Student ID","Name","Age","Gender","Course","Email","Phone"])
        users = wb.create_sheet("Users")
        users.append(["Username","Password"])
        wb.save(FILE)
        wb.close()

def register_window():
    win = tk.Toplevel(root)
    win.title("Register")
    win.geometry("350x300")
    tk.Label(win,text="Register User",font=("Arial",18,"bold")).pack(pady=15)

    tk.Label(win,text="Username").pack()
    username = tk.Entry(win,width=30)
    username.pack(pady=5)

    tk.Label(win,text="Password").pack()
    password = tk.Entry(win,width=30,show="*")
    password.pack(pady=5)

    tk.Label(win,text="Confirm Password").pack()
    confirm = tk.Entry(win,width=30,show="*")
    confirm.pack(pady=5)

    def register():
        u,p,c = username.get().strip(),password.get(),confirm.get()

        if not u or not p:
            messagebox.showwarning("Warning","Fill all fields")
            return
        if p != c:
            messagebox.showerror("Error","Passwords do not match")
            return

        wb = load_workbook(FILE)
        ws = wb["Users"]

        for row in ws.iter_rows(min_row=2,values_only=True):
            if row[0] == u:
                messagebox.showerror("Error","Username already exists")
                wb.close()
                return

        ws.append([u,hash_password(p)])
        wb.save(FILE)
        wb.close()
        messagebox.showinfo("Success","Registration successful")
        win.destroy()

    tk.Button(win,text="Register",command=register,width=15).pack(pady=15)

def clear():
    for entry in entries:
        entry.delete(0,tk.END)
    gender.set("")
    course.set("")

def get_data():
    return [
        entries[0].get().strip(),
        entries[1].get().strip(),
        entries[2].get().strip(),
        gender.get(),
        course.get(),
        entries[3].get().strip(),
        entries[4].get().strip()
    ]

def show_students():
    for item in table.get_children():
        table.delete(item)

    wb = load_workbook(FILE)
    ws = wb["Students"]

    for row in ws.iter_rows(min_row=2,values_only=True):
        table.insert("",tk.END,values=row)

    wb.close()

def add_student():
    data = get_data()

    if not data[0] or not data[1]:
        messagebox.showwarning("Warning","Student ID and Name required")
        return

    wb = load_workbook(FILE)
    ws = wb["Students"]

    for row in ws.iter_rows(min_row=2,values_only=True):
        if str(row[0]) == data[0]:
            messagebox.showerror("Error","Student ID already exists")
            wb.close()
            return

    ws.append(data)
    wb.save(FILE)
    wb.close()
    messagebox.showinfo("Success","Student added successfully")
    clear()
    show_students()

def select_student(event):
    selected = table.focus()
    if not selected:
        return

    values = table.item(selected,"values")
    clear()

    for i in range(3):
        entries[i].insert(0,values[i])

    gender.set(values[3])
    course.set(values[4])
    entries[3].insert(0,values[5])
    entries[4].insert(0,values[6])

def update_student():
    data = get_data()

    if not data[0]:
        messagebox.showwarning("Warning","Enter Student ID")
        return

    wb = load_workbook(FILE)
    ws = wb["Students"]
    found = False

    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == data[0]:
            for i,value in enumerate(data,1):
                row[i-1].value = value
            found = True
            break

    if found:
        wb.save(FILE)
        messagebox.showinfo("Success","Student updated")
    else:
        messagebox.showerror("Error","Student not found")

    wb.close()
    show_students()

def delete_student():
    student_id = entries[0].get().strip()

    if not student_id:
        messagebox.showwarning("Warning","Enter Student ID")
        return

    wb = load_workbook(FILE)
    ws = wb["Students"]
    found = False

    for row in range(2,ws.max_row+1):
        if str(ws.cell(row,1).value) == student_id:
            ws.delete_rows(row,1)
            found = True
            break

    if found:
        wb.save(FILE)
        messagebox.showinfo("Success","Student deleted")
    else:
        messagebox.showerror("Error","Student not found")

    wb.close()
    clear()
    show_students()

def search_student():
    keyword = search_entry.get().lower()

    for item in table.get_children():
        table.delete(item)

    wb = load_workbook(FILE)
    ws = wb["Students"]

    for row in ws.iter_rows(min_row=2,values_only=True):
        if any(keyword in str(value).lower() for value in row):
            table.insert("",tk.END,values=row)

    wb.close()

def login():
    u = username_entry.get().strip()
    p = password_entry.get()

    wb = load_workbook(FILE)
    ws = wb["Users"]
    valid = False

    for row in ws.iter_rows(min_row=2,values_only=True):
        if row[0] == u and row[1] == hash_password(p):
            valid = True
            break

    wb.close()

    if valid:
        login_frame.pack_forget()
        student_frame.pack(fill="both",expand=True)
        show_students()
    else:
        messagebox.showerror("Login Failed","Invalid username or password")

def logout():
    student_frame.pack_forget()
    login_frame.pack(fill="both",expand=True)
    username_entry.delete(0,tk.END)
    password_entry.delete(0,tk.END)
    clear()

create_file()

root = tk.Tk()
root.title("Student Management System")
root.geometry("900x600")

login_frame = tk.Frame(root)
login_frame.pack(fill="both",expand=True)

tk.Label(login_frame,text="Student Management System",
         font=("Arial",22,"bold")).pack(pady=40)

tk.Label(login_frame,text="Username").pack()
username_entry = tk.Entry(login_frame,width=30)
username_entry.pack(pady=5)

tk.Label(login_frame,text="Password").pack()
password_entry = tk.Entry(login_frame,width=30,show="*")
password_entry.pack(pady=5)

tk.Button(login_frame,text="Login",command=login,width=15).pack(pady=10)
tk.Button(login_frame,text="Register",
          command=register_window,width=15).pack()

student_frame = tk.Frame(root)

tk.Label(student_frame,text="Student Management System",
         font=("Arial",20,"bold")).pack(pady=10)

form = tk.Frame(student_frame)
form.pack()

labels = ["Student ID","Name","Age"]
entries = []

for i,label in enumerate(labels):
    tk.Label(form,text=label).grid(row=i,column=0,padx=5,pady=5)
    e = tk.Entry(form,width=25)
    e.grid(row=i,column=1,padx=5,pady=5)
    entries.append(e)

tk.Label(form,text="Gender").grid(row=0,column=2)
gender = ttk.Combobox(form,values=["Male","Female","Other"],
                      width=22,state="readonly")
gender.grid(row=0,column=3,padx=5)

tk.Label(form,text="Course").grid(row=1,column=2)
course = ttk.Combobox(form,
                      values=["Computer","IT","Mechanical","Civil","Electrical"],
                      width=22)
course.grid(row=1,column=3,padx=5)

tk.Label(form,text="Email").grid(row=2,column=2)
email = tk.Entry(form,width=25)
email.grid(row=2,column=3,padx=5)
entries.append(email)

tk.Label(form,text="Phone").grid(row=3,column=2)
phone = tk.Entry(form,width=25)
phone.grid(row=3,column=3,padx=5)
entries.append(phone)

buttons = tk.Frame(student_frame)
buttons.pack(pady=10)

for i,(text,command) in enumerate([
    ("Add",add_student),
    ("Update",update_student),
    ("Delete",delete_student),
    ("Clear",clear),
    ("Logout",logout)
]):
    tk.Button(buttons,text=text,command=command,width=12).grid(
        row=0,column=i,padx=3)

search_entry = tk.Entry(student_frame,width=30)
search_entry.pack(pady=5)

tk.Button(student_frame,text="Search",
          command=search_student).pack()

columns = ["Student ID","Name","Age","Gender","Course","Email","Phone"]

table = ttk.Treeview(student_frame,columns=columns,
                     show="headings",height=10)

for col in columns:
    table.heading(col,text=col)
    table.column(col,width=110)

table.pack(fill="both",expand=True,padx=10,pady=10)
table.bind("<ButtonRelease-1>",select_student)

student_frame.pack_forget()
root.mainloop()